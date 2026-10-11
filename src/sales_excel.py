"""Builds the formatted sales analysis Excel workbook from the computed KPIs."""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .clean import QualityLog
from .excel_style import (
    BOLD,
    add_table_polish,
    write_dataframe,
    write_findings,
    write_header_row,
    write_title,
)
from .sales_kpis import SalesReportData

COUNT_FORMAT = "#,##0"


def _currency_format(symbol: str) -> str:
    return f'"{symbol}"#,##0.00' if symbol else "#,##0.00"


def _revenue_label(symbol: str) -> str:
    return f"Revenue ({symbol})" if symbol else "Revenue"

# Distinct tab colors so the sheets are easy to tell apart at a glance.
TAB_COLORS = {
    "Summary": "1F4E78",
    "Top Products": "2E7D32",
    "By Country": "6A1B9A",
    "Data Quality Log": "757575",
}


def _autofit_first_column(ws: Worksheet, width: int = 22) -> None:
    ws.column_dimensions["A"].width = width


def _build_summary_sheet(ws: Worksheet, data: SalesReportData) -> None:
    write_title(ws, f"Sales Analysis: {data.period_label}", span_cols=4)
    _autofit_first_column(ws, 26)
    currency_format = _currency_format(data.currency)
    t = data.totals
    cancellations_value_label = f"Cancellations (value, {data.currency})" if data.currency else "Cancellations (value)"

    # (label, value, format kind)
    kpi_rows = [
        (_revenue_label(data.currency), t["revenue"], "currency"),
        ("Orders", t["orders"], "count"),
        ("Average Order Value", t["avg_order_value"], "currency"),
        ("Units Sold", t["units"], "count"),
        ("Unique Customers", t["customers"], "count"),
        ("Cancelled Orders", t["cancellations_count"], "count"),
        (cancellations_value_label, t["cancellations_value"], "currency"),
    ]
    kpi_df = pd.DataFrame([row[:2] for row in kpi_rows], columns=["Metric", "Value"])

    header_row = 3
    last_row = write_dataframe(ws, kpi_df, header_row)
    for offset, (_, value, kind) in enumerate(kpi_rows):
        cell = ws.cell(row=header_row + 1 + offset, column=2)
        if value is None:  # e.g. customers, when the sheet has no customer column
            cell.value = "n/a"
            cell.alignment = Alignment(horizontal="center")
        else:
            cell.number_format = currency_format if kind == "currency" else COUNT_FORMAT
    ws.column_dimensions["B"].width = 18
    ws.freeze_panes = f"A{header_row + 1}"

    last_row = write_findings(ws, data.findings, start_row=last_row + 2)

    # Month-by-month table (feeds the chart beside the KPI table)
    table_row = last_row + 3
    ws.cell(row=table_row, column=1, value="Monthly Revenue").font = BOLD
    monthly_last_row = write_dataframe(
        ws, data.monthly, table_row + 1, number_formats={0: "mmm yyyy", 1: currency_format, 2: COUNT_FORMAT}
    )

    chart = LineChart()
    chart.title = "Monthly Revenue"
    chart.y_axis.title = _revenue_label(data.currency)
    chart.x_axis.number_format = "mmm yyyy"
    values = Reference(ws, min_col=2, min_row=table_row + 1, max_row=monthly_last_row)
    categories = Reference(ws, min_col=1, min_row=table_row + 2, max_row=monthly_last_row)
    chart.add_data(values, titles_from_data=True)
    chart.set_categories(categories)
    chart.width = 20
    chart.height = 10
    ws.add_chart(chart, f"F{header_row}")


def _build_top_products_sheet(ws: Worksheet, data: SalesReportData) -> None:
    write_title(ws, "Top 10 Products by Revenue", span_cols=3)
    header_row = 3
    last_row = write_dataframe(
        ws, data.top_products, header_row, number_formats={2: _currency_format(data.currency), 3: COUNT_FORMAT}
    )
    add_table_polish(ws, header_row, last_row, last_col=len(data.top_products.columns))

    chart = BarChart()
    chart.title = "Top Products by Revenue"
    chart.y_axis.title = _revenue_label(data.currency)
    values = Reference(ws, min_col=3, min_row=header_row, max_row=last_row)
    categories = Reference(ws, min_col=2, min_row=header_row + 1, max_row=last_row)
    chart.add_data(values, titles_from_data=True)
    chart.set_categories(categories)
    chart.width = 22
    chart.height = 12
    ws.add_chart(chart, "F3")


def _build_by_country_sheet(ws: Worksheet, data: SalesReportData) -> None:
    write_title(ws, "Revenue by Country", span_cols=3)
    header_row = 3
    last_row = write_dataframe(
        ws, data.by_country, header_row, number_formats={1: _currency_format(data.currency), 2: COUNT_FORMAT}
    )
    add_table_polish(ws, header_row, last_row, last_col=len(data.by_country.columns))

    chart = BarChart()
    chart.title = "Revenue by Country"
    chart.y_axis.title = _revenue_label(data.currency)
    values = Reference(ws, min_col=2, min_row=header_row, max_row=last_row)
    categories = Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_row)
    chart.add_data(values, titles_from_data=True)
    chart.set_categories(categories)
    chart.width = 22
    chart.height = 12
    ws.add_chart(chart, "E3")


def _build_quality_log_sheet(ws: Worksheet, quality_log: QualityLog) -> None:
    write_title(ws, "Data Quality Log", span_cols=3)
    _autofit_first_column(ws, 30)
    descriptions = {
        "rows_in": "Raw rows read from source export",
        "duplicates_dropped": "Exact duplicate rows removed",
        "unreadable_rows_dropped": "Rows dropped because the date, quantity or price couldn't be read",
        "cancellations_separated": "Cancelled-order lines separated into returns reporting",
        "missing_customer_id": "Sales rows with a missing CustomerID (kept, revenue still counted)",
        "non_product_dropped": "Postage, fees and other non-product lines dropped",
        "non_positive_price_dropped": "Rows dropped for a zero or negative price or quantity",
        "rows_out": "Clean product-sale rows used for KPI calculations",
    }
    df = pd.DataFrame(
        [(key.replace("_", " ").title(), value, descriptions[key]) for key, value in quality_log._asdict().items()],
        columns=["Check", "Count", "Description"],
    )
    header_row = 3
    last_row = write_dataframe(ws, df, header_row, number_formats={1: COUNT_FORMAT})
    add_table_polish(ws, header_row, last_row, last_col=len(df.columns))


ORGANIZED_COLUMNS = {
    "InvoiceDate": ("Date", 18),
    "InvoiceNo": ("Order ID", 14),
    "StockCode": ("Product Code", 16),
    "Description": ("Product", 40),
    "Quantity": ("Quantity", 10),
    "UnitPrice": ("Unit Price", 12),
    "LineTotal": ("Line Total", 13),
    "CustomerID": ("Customer", 14),
    "Country": ("Country", 18),
}
# Above this many rows, only the most recent ones are written, to keep the file a sensible size.
ORGANIZED_ROW_LIMIT = 100_000


def _build_organized_sheet(ws: Worksheet, sales_df: pd.DataFrame, data: SalesReportData) -> None:
    rows = sales_df.sort_values(["InvoiceDate", "InvoiceNo"])
    if len(rows) > ORGANIZED_ROW_LIMIT:
        rows = rows.tail(ORGANIZED_ROW_LIMIT)
        title = f"Organized Data: the latest {len(rows):,} of {len(sales_df):,} clean sales rows"
    else:
        title = f"Organized Data: all {len(rows):,} clean sales rows"

    rows = rows[list(ORGANIZED_COLUMNS)]
    rows = rows.astype(object).where(rows.notna(), None)

    write_title(ws, title, span_cols=len(ORGANIZED_COLUMNS))
    header_row = 3
    write_header_row(ws, [label for label, _ in ORGANIZED_COLUMNS.values()], header_row)
    for values in rows.itertuples(index=False):
        ws.append(list(values))

    last_row = header_row + len(rows)
    formats = {1: "yyyy-mm-dd hh:mm", 6: _currency_format(data.currency), 7: _currency_format(data.currency), 8: "0"}
    for col, number_format in formats.items():
        for (cell,) in ws.iter_rows(min_row=header_row + 1, max_row=last_row, min_col=col, max_col=col):
            cell.number_format = number_format
    for col, (_, width) in enumerate(ORGANIZED_COLUMNS.values(), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    add_table_polish(ws, header_row, max(last_row, header_row), last_col=len(ORGANIZED_COLUMNS))
    ws.sheet_properties.tabColor = "EF6C00"


def build_workbook(
    data: SalesReportData, quality_log: QualityLog, out_path: Path, organized_rows: pd.DataFrame | None = None
) -> None:
    """`organized_rows` (the cleaned sales rows), if given, adds an "Organized Data" sheet."""
    wb = Workbook()

    summary_ws = wb.active
    summary_ws.title = "Summary"
    _build_summary_sheet(summary_ws, data)

    _build_top_products_sheet(wb.create_sheet("Top Products"), data)
    _build_by_country_sheet(wb.create_sheet("By Country"), data)
    if organized_rows is not None:
        _build_organized_sheet(wb.create_sheet("Organized Data"), organized_rows, data)
    _build_quality_log_sheet(wb.create_sheet("Data Quality Log"), quality_log)

    for sheet_name, color in TAB_COLORS.items():
        wb[sheet_name].sheet_properties.tabColor = color

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
