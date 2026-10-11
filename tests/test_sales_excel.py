import dataclasses

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.clean import QualityLog
from src.sales_excel import TAB_COLORS, build_workbook
from src.sales_kpis import build_report_data

EXPECTED_SHEETS = ["Summary", "Top Products", "By Country", "Data Quality Log"]


def _sales_row(invoice_no, stock_code, description, quantity, invoice_date, unit_price, customer_id, country):
    return {
        "InvoiceNo": invoice_no,
        "StockCode": stock_code,
        "Description": description,
        "Quantity": quantity,
        "InvoiceDate": pd.Timestamp(invoice_date),
        "UnitPrice": unit_price,
        "CustomerID": customer_id,
        "Country": country,
        "LineTotal": quantity * unit_price,
    }


@pytest.fixture
def quality_log():
    return QualityLog(
        rows_in=100,
        duplicates_dropped=2,
        unreadable_rows_dropped=1,
        cancellations_separated=3,
        missing_customer_id=1,
        non_product_dropped=0,
        non_positive_price_dropped=4,
        rows_out=90,
    )


@pytest.fixture
def report_data():
    """A few sales across two weeks of November, plus one cancellation."""
    sales_df = pd.DataFrame(
        [
            _sales_row("600001", "A1", "Widget", 10, "2011-11-21 10:00:00", 10.0, 1001, "United Kingdom"),
            _sales_row("600002", "B1", "Gadget", 5, "2011-11-23 11:00:00", 10.0, 1002, "France"),
            _sales_row("500001", "A1", "Widget", 10, "2011-11-14 10:00:00", 10.0, 2001, "United Kingdom"),
        ]
    )
    cancellations_df = pd.DataFrame(
        [_sales_row("C600004", "A1", "Widget", -2, "2011-11-22 08:00:00", 10.0, 1001, "United Kingdom")]
    )
    return build_report_data(sales_df, cancellations_df)


@pytest.fixture
def report_data_single_row():
    """The smallest possible dataset: one sale and no cancellations."""
    sales_df = pd.DataFrame(
        [_sales_row("600001", "A1", "Widget", 10, "2011-11-21 10:00:00", 10.0, 1001, "United Kingdom")]
    )
    cancellations_df = pd.DataFrame(columns=list(sales_df.columns))
    return build_report_data(sales_df, cancellations_df)


class TestBuildWorkbook:
    def test_creates_all_expected_sheets(self, report_data, quality_log, tmp_path):
        out_path = tmp_path / "report.xlsx"
        build_workbook(report_data, quality_log, out_path)

        assert out_path.exists()
        wb = load_workbook(out_path)
        assert wb.sheetnames == EXPECTED_SHEETS

    def test_sheet_tabs_are_colored(self, report_data, quality_log, tmp_path):
        out_path = tmp_path / "report.xlsx"
        build_workbook(report_data, quality_log, out_path)

        wb = load_workbook(out_path)
        for sheet_name, expected_color in TAB_COLORS.items():
            assert wb[sheet_name].sheet_properties.tabColor.rgb.endswith(expected_color)

    def test_data_sheets_have_frozen_header_and_autofilter(self, report_data, quality_log, tmp_path):
        out_path = tmp_path / "report.xlsx"
        build_workbook(report_data, quality_log, out_path)

        wb = load_workbook(out_path)
        for sheet_name in ("Top Products", "By Country", "Data Quality Log"):
            ws = wb[sheet_name]
            assert ws.freeze_panes == "A4"
            assert ws.auto_filter.ref is not None

    def test_customers_show_n_a_without_customer_ids(self, quality_log, tmp_path):
        sales_df = pd.DataFrame(
            [
                _sales_row("600001", "A1", "Widget", 10, "2011-11-21 10:00:00", 10.0, None, "United Kingdom"),
                _sales_row("500001", "A1", "Widget", 10, "2011-11-14 10:00:00", 10.0, None, "United Kingdom"),
            ]
        )
        data = build_report_data(sales_df, sales_df.iloc[0:0])
        out_path = tmp_path / "report.xlsx"
        build_workbook(data, quality_log, out_path)

        ws = load_workbook(out_path)["Summary"]
        assert ws["A8"].value == "Unique Customers"
        assert ws["B8"].value == "n/a"

    def test_uses_the_report_currency_in_labels_and_formats(self, report_data, quality_log, tmp_path):
        out_path = tmp_path / "report.xlsx"
        build_workbook(dataclasses.replace(report_data, currency="$"), quality_log, out_path)

        ws = load_workbook(out_path)["Summary"]
        assert ws["A4"].value == "Revenue ($)"
        assert ws["B4"].number_format == '"$"#,##0.00'

    def test_summary_shows_key_findings_under_the_kpi_table(self, report_data, quality_log, tmp_path):
        report_data.findings = ["The best month was November 2011.", "Monday was the busiest day of the week."]
        out_path = tmp_path / "report.xlsx"
        build_workbook(report_data, quality_log, out_path)

        ws = load_workbook(out_path)["Summary"]
        assert (ws["A4"].value, ws["A5"].value) == ("Revenue", "Orders")
        assert ws["A12"].value == "Key findings"
        assert ws["A13"].value == "• The best month was November 2011."
        assert ws["A14"].value == "• Monday was the busiest day of the week."
        assert "A13:D13" in {str(r) for r in ws.merged_cells.ranges}
        assert ws["A17"].value == "Monthly Revenue"
        assert ws["A19"].number_format == "mmm yyyy"

    def test_organized_data_sheet_lists_clean_rows_sorted_by_date(self, report_data, quality_log, tmp_path):
        rows = pd.DataFrame(
            [
                _sales_row("600002", "B1", "Gadget", 5, "2011-11-23 11:00:00", 10.0, None, "France"),
                _sales_row("600001", "A1", "Widget", 10, "2011-11-21 10:00:00", 10.0, 1001, "United Kingdom"),
            ]
        )
        out_path = tmp_path / "report.xlsx"
        build_workbook(report_data, quality_log, out_path, organized_rows=rows)

        wb = load_workbook(out_path)
        assert wb.sheetnames == ["Summary", "Top Products", "By Country", "Organized Data", "Data Quality Log"]
        ws = wb["Organized Data"]
        assert [c.value for c in ws[3]] == [
            "Date", "Order ID", "Product Code", "Product", "Quantity", "Unit Price", "Line Total", "Customer", "Country"
        ]
        assert ws["B4"].value == "600001" and ws["B5"].value == "600002"
        assert ws["H5"].value is None
        assert ws.freeze_panes == "A4"

    def test_organized_data_keeps_the_latest_rows_for_large_data(self, report_data, quality_log, tmp_path, monkeypatch):
        monkeypatch.setattr("src.sales_excel.ORGANIZED_ROW_LIMIT", 2)
        rows = pd.DataFrame(
            [
                _sales_row("600001", "A1", "Widget", 10, "2011-11-21 10:00:00", 10.0, 1001, "UK"),
                _sales_row("500001", "A1", "Widget", 10, "2011-11-14 10:00:00", 10.0, 2001, "UK"),
                _sales_row("100001", "A1", "Widget", 10, "2011-01-03 10:00:00", 10.0, 3001, "UK"),
            ]
        )
        out_path = tmp_path / "report.xlsx"
        build_workbook(report_data, quality_log, out_path, organized_rows=rows)

        ws = load_workbook(out_path)["Organized Data"]
        assert [ws["B4"].value, ws["B5"].value, ws["B6"].value] == ["500001", "600001", None]
        assert "the latest 2 of 3" in ws["A1"].value

    def test_creates_output_directory_if_missing(self, report_data, quality_log, tmp_path):
        out_path = tmp_path / "nested" / "dir" / "report.xlsx"
        build_workbook(report_data, quality_log, out_path)
        assert out_path.exists()

    def test_single_row_data_still_builds_every_sheet(self, report_data_single_row, quality_log, tmp_path):
        out_path = tmp_path / "report.xlsx"
        build_workbook(report_data_single_row, quality_log, out_path)
        wb = load_workbook(out_path)
        assert wb["Top Products"]["A3"].value == "StockCode"
        assert wb["By Country"]["A3"].value == "Country"
